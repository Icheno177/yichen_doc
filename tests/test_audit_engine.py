from __future__ import annotations

import json
from pathlib import Path
import tempfile
import unittest

from PIL import Image
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from yichen_doc.audit.content_mapping import map_document_windows, map_units, mapping_stats
from yichen_doc.audit.conflict_detection import reconcile_checks
from yichen_doc.audit.evidence_audit import audit_claims
from yichen_doc.audit.extractors import extract
from yichen_doc.audit.model_audit import audit_models, evaluate_expression
from yichen_doc.audit.schemas import ContentUnit, MappingStatus, sha256_file, unit_id
from yichen_doc.audit.visual_audit import audit_rendered_pages
from yichen_doc.audit.format_audits import audit_html
from yichen_doc.audit.format_audits import audit_docx


def unit(name: str, locator: str, text: str, kind: str = "paragraph") -> ContentUnit:
    digest = (name.encode().hex() + "0" * 64)[:64]
    return ContentUnit(unit_id(digest, locator, text), name, digest, locator, "test", kind, text)


class MappingTests(unittest.TestCase):
    def test_01_exact_mapping(self):
        s, t = unit("s", "p1", "Revenue grew 20%"), unit("t", "p1", "Revenue grew 20%")
        self.assertEqual(map_units([s], [t])[0].mapping_status, MappingStatus.RETAINED)

    def test_02_whitespace_normalized_mapping(self):
        r = map_units([unit("s", "p1", "revenue grew")], [unit("t", "p1", "revenue   grew")])[0]
        self.assertEqual(r.mapping_status, MappingStatus.RETAINED)

    def test_03_punctuation_rewrite_requires_review(self):
        r = map_units([unit("s", "p1", "A, B")], [unit("t", "p1", "A B")])[0]
        self.assertEqual(r.mapping_status, MappingStatus.REWRITTEN)
        self.assertEqual(r.reviewer_status, "REVIEW_REQUIRED")

    def test_04_source_split_across_target_units(self):
        r = map_units([unit("s", "p1", "alpha beta")], [unit("t", "l1", "alpha"), unit("t", "l2", "beta")])[0]
        self.assertEqual(r.mapping_status, MappingStatus.SPLIT)

    def test_05_source_absence_is_unresolved_not_confirmed_missing(self):
        r = map_units([unit("s", "p1", "a long source unit absent")], [unit("t", "p1", "unrelated text")])[0]
        self.assertEqual(r.mapping_status, MappingStatus.UNRESOLVED)

    def test_06_reason_without_authorization_does_not_remove_source(self):
        s = unit("s", "p1", "source sentence")
        r = map_units([s], [unit("t", "p1", "other text")], authorized_removals={s.unit_id: "User authorized removal"})[0]
        self.assertEqual(r.mapping_status, MappingStatus.UNRESOLVED)

    def test_35_authorized_removal_records_reviewer_and_basis(self):
        s = unit("s", "p1", "source sentence")
        decision = {"reason": "Out of agreed scope", "authorization": "User request 2026-10-26", "reviewer": "Analyst", "reviewed_at": "2026-10-26"}
        r = map_units([s], [unit("t", "p1", "other text")], authorized_removals={s.unit_id: decision})[0]
        self.assertEqual(r.mapping_status, MappingStatus.INTENTIONALLY_REMOVED)
        self.assertEqual(r.reviewer_status, "VERIFIED")

    def test_07_numeric_conflict_detected(self):
        r = map_units([unit("s", "p1", "revenue grew 20% in 2025")], [unit("t", "p1", "revenue grew 10% in 2025")])[0]
        self.assertEqual(r.mapping_status, MappingStatus.CONFLICT)
        self.assertEqual(r.numeric_consistency, "CONFLICT")

    def test_08_negation_conflict_detected(self):
        r = map_units([unit("s", "p1", "the device is approved")], [unit("t", "p1", "the device is not approved")])[0]
        self.assertEqual(r.mapping_status, MappingStatus.CONFLICT)

    def test_09_direction_conflict_detected(self):
        r = map_units([unit("s", "p1", "revenue grew 20%")], [unit("t", "p1", "revenue fell 20%")])[0]
        self.assertEqual(r.mapping_status, MappingStatus.CONFLICT)

    def test_10_lexical_similarity_is_candidate_only(self):
        r = map_units([unit("s", "p1", "commercialization economics and customer renewal profile")], [unit("t", "p1", "commercial economics; customer renewal profiles")])[0]
        self.assertIn(r.mapping_status, {MappingStatus.UNRESOLVED, MappingStatus.CONFLICT})
        self.assertEqual(r.reviewer_status, "REVIEW_REQUIRED")

    def test_36_legacy_document_window_screen_is_separate_and_reproducible(self):
        source = [unit("s", "p1", "a" * 80 + "b" * 86)]
        target = [unit("t", "p1", "a" * 80)]
        stats = mapping_stats([], source, target)
        legacy = stats["legacy_document_wide_window_screen"]
        self.assertEqual(legacy["denominator"], 2)
        self.assertEqual(legacy["exact_matches"], 1)
        self.assertIn("do not prove semantic retention", legacy["interpretation"])
        rows = map_document_windows(source, target)
        self.assertEqual([x["mapping_status"] for x in rows], ["RETAINED", "UNRESOLVED"])
        self.assertTrue(rows[1]["source_locator"])

    def test_43_destination_convergence_is_not_auto_labeled_merge(self):
        a = unit("s", "p1", "Repeated source statement")
        b = unit("s", "p2", "Repeated source statement")
        target = unit("t", "slide1", "Repeated source statement")
        self.assertEqual([r.mapping_status for r in map_units([a, b], [target])], [MappingStatus.RETAINED, MappingStatus.RETAINED])


class ModelTests(unittest.TestCase):
    def test_11_roi_recalculation(self):
        r = audit_models([{"model_id":"R1","model_type":"roi","inputs":{"incremental_benefit":{"value":150,"classification":"VERIFIED_INTERNAL"},"incremental_cost":{"value":60,"classification":"VERIFIED_INTERNAL"}}}])
        self.assertAlmostEqual(r["models"][0]["independent_output"], 1.5)

    def test_12_npv_recalculation(self):
        r = audit_models([{"model_type":"npv","cashflows":[-100,60,60],"discount_rate":0.1}])
        self.assertAlmostEqual(r["models"][0]["independent_output"], 4.1322314049)

    def test_13_irr_recalculation(self):
        r = audit_models([{"model_type":"irr","cashflows":[-100,60,60]}])
        self.assertAlmostEqual(r["models"][0]["independent_output"], 0.1306623863, places=6)

    def test_14_payback_interpolates_period(self):
        r = audit_models([{"model_type":"payback","cashflows":[-100,60,60]}])
        self.assertAlmostEqual(r["models"][0]["independent_output"], 1.6666666667)

    def test_15_missing_input_is_open_issue(self):
        r = audit_models([{"model_type":"roi","inputs":{"incremental_benefit":{"value":None,"classification":"MISSING"},"incremental_cost":{"value":20,"classification":"VERIFIED_INTERNAL"}}}])
        self.assertTrue(any("MODEL-MISSING" in x["issue_id"] for x in r["issues"]))

    def test_16_unverified_input_stays_conditional(self):
        r = audit_models([{"model_type":"roi","inputs":{"incremental_benefit":{"value":100,"classification":"ASSUMPTION"},"incremental_cost":{"value":50,"classification":"ASSUMPTION"}}}])
        self.assertEqual(r["models"][0]["result_status"], "RECALCULATED_CONDITIONAL")
        self.assertTrue(any("MODEL-UNVERIFIED" in x["issue_id"] for x in r["issues"]))

    def test_17_reported_output_mismatch_is_critical(self):
        r = audit_models([{"model_type":"roi","reported_output":9,"inputs":{"incremental_benefit":{"value":100,"classification":"VERIFIED_INTERNAL"},"incremental_cost":{"value":50,"classification":"VERIFIED_INTERNAL"}}}])
        self.assertTrue(any(x["severity"] == "CRITICAL" for x in r["issues"]))

    def test_18_expression_evaluator_rejects_function_calls(self):
        with self.assertRaises(ValueError): evaluate_expression("__import__('os').system('x')", {})

    def test_19_expression_divide_by_zero_is_rejected_by_model_runner(self):
        r = audit_models([{"model_id":"E1","expression":"a / b","inputs":{"a":{"value":2,"classification":"ASSUMPTION"},"b":{"value":0,"classification":"ASSUMPTION"}}}])
        self.assertIsNone(r["models"][0]["independent_output"])

    def test_38_tco_aggregates_classified_costs(self):
        r = audit_models([{"model_type":"tco","inputs":{"license":{"value":10,"classification":"ASSUMPTION"},"cloud":{"value":5,"classification":"ASSUMPTION"}}}])
        self.assertEqual(r["models"][0]["independent_output"], 15)

    def test_39_cagr_recalculation(self):
        r = audit_models([{"model_type":"cagr","periods":2,"inputs":{"start":{"value":100,"classification":"ASSUMPTION"},"end":{"value":121,"classification":"ASSUMPTION"}}}])
        self.assertAlmostEqual(r["models"][0]["independent_output"], 0.1)

    def test_40_dcf_terminal_value_is_discounted(self):
        r = audit_models([{"model_type":"dcf","cashflows":[0,10],"discount_rate":0.1,"terminal_value":100}])
        self.assertAlmostEqual(r["models"][0]["independent_output"], 100)

    def test_41_scenarios_and_sensitivity_are_recalculated(self):
        models = [
            {"model_type":"scenarios","expression":"price * volume","scenarios":[{"name":"Bear","inputs":{"price":2,"volume":3}},{"name":"Bull","inputs":{"price":4,"volume":3}}]},
            {"model_type":"sensitivity","expression":"price * volume","inputs":{"volume":{"value":3,"classification":"ASSUMPTION"}},"parameter":"price","values":[2,4]},
        ]
        r = audit_models(models)
        self.assertEqual([x["independent_output"] for x in r["models"][0]["independent_output"]], [6,12])
        self.assertEqual([x["independent_output"] for x in r["models"][1]["independent_output"]], [6,12])

    def test_42_gmv_tam_sam_som_expression_models(self):
        r = audit_models([
            {"model_type":"gmv","inputs":{"transaction_count":{"value":10,"classification":"ASSUMPTION"},"average_order_value":{"value":5,"classification":"ASSUMPTION"}}},
            {"model_type":"tam","expression":"eligible_customers * annual_spend","inputs":{"eligible_customers":{"value":100,"classification":"ASSUMPTION"},"annual_spend":{"value":20,"classification":"ASSUMPTION"}}},
        ])
        self.assertEqual(r["models"][0]["independent_output"], 50)
        self.assertEqual(r["models"][1]["independent_output"], 2000)


class EvidenceAndScopeTests(unittest.TestCase):
    def test_20_unknown_evidence_id_reported(self):
        r = audit_claims([{"claim_id":"C1","text":"claim","label":"FACT","source_ids":["S2"]}], [{"id":"S1","title":"source","grade":"A"}])
        self.assertTrue(any("UNKNOWN" in x["description"].upper() for x in r["issues"]))

    def test_21_invalid_evidence_grade_reported(self):
        r = audit_claims([{"claim_id":"C1","text":"claim","label":"FACT","source_ids":["S1"]}], [{"id":"S1","title":"source","grade":"Z"}])
        self.assertTrue(any("grade" in x["description"].lower() for x in r["issues"]))

    def test_22_url_not_checked_means_unverified(self):
        r = audit_claims([{"claim_id":"C1","text":"claim","label":"FACT","source_ids":["S1"]}], [{"id":"S1","title":"source","grade":"A","url":"https://example.com"}])
        self.assertEqual(r["ledger"][0]["verification_status"], "UNVERIFIED")

    def test_23_same_scope_results_conflict(self):
        r = reconcile_checks([{"check_id":"QA-1","check_name":"coverage","input_scope":"A→B","denominator":"10","method":"exact","value":0.8},{"check_id":"QA-1","check_name":"coverage","input_scope":"A→B","denominator":"10","method":"exact","value":0.9}])
        self.assertEqual(len(r["conflicts"]), 1)

    def test_24_different_denominator_is_non_comparable(self):
        r = reconcile_checks([{"check_id":"QA-1","check_name":"coverage","input_scope":"A→B","denominator":"10","method":"exact","value":0.8},{"check_id":"QA-1","check_name":"coverage","input_scope":"A→B","denominator":"20","method":"exact","value":0.9}])
        self.assertEqual(len(r["conflicts"]), 0)
        self.assertEqual(len(r["non_comparable"]), 1)


class ExtractorAndVisualTests(unittest.TestCase):
    def test_25_docx_paragraph_and_table_cells_extracted(self):
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "x.docx"; doc = Document(); doc.add_paragraph("claim"); t = doc.add_table(rows=1, cols=1); t.cell(0,0).text = "cell"; doc.save(p)
            kinds = {x.kind for x in extract(p)}
            self.assertIn("paragraph", kinds); self.assertIn("table_cell", kinds)
            audit = audit_docx(p, Path("golden_sets/word/style_manifest.json"))
            self.assertEqual(audit["render_status"], "NOT_RUN")
            self.assertIn("golden_style_structure_comparison", audit)

    def test_26_pptx_text_extracted(self):
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "x.pptx"; prs = Presentation(); s = prs.slides.add_slide(prs.slide_layouts[6]); s.shapes.add_textbox(0,0,1000000,1000000).text = "slide text"; prs.save(p)
            self.assertTrue(any("slide text" in x.text for x in extract(p)))

    def test_27_xlsx_formula_extracted(self):
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "x.xlsx"; wb = Workbook(); wb.active["A1"] = "=1+2"; wb.save(p)
            self.assertEqual(next(x for x in extract(p) if x.kind == "formula").text, "=1+2")

    def test_28_html_local_link_and_text_extracted(self):
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "x.html"; p.write_text('<h1>Title</h1><p>Paragraph</p><a href="https://example.com">Source</a>', encoding="utf-8")
            units = extract(p)
            self.assertTrue(any(x.kind == "heading" for x in units)); self.assertTrue(any(x.kind == "citation" for x in units))

    def test_33_nested_html_inline_text_is_not_lost(self):
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "nested.html"; p.write_text('<p>Revenue <strong>grew</strong> 20%.</p>', encoding="utf-8")
            self.assertTrue(any(x.text == "Revenue grew 20%." for x in extract(p)))

    def test_34_short_entity_occurrence_is_not_claimed_as_retained(self):
        source = unit("s", "table/cell1", "企业")
        target = [unit("t", f"line/{i}", f"企业研究第{i}章") for i in range(1, 5)]
        record = map_units([source], target)[0]
        self.assertEqual(record.mapping_status, MappingStatus.UNRESOLVED)

    def test_37_html_browser_audit_never_writes_into_source_directory(self):
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "report.html"; p.write_text("<title>Report</title>", encoding="utf-8")
            result = audit_html(p, browser_check=True)
            self.assertEqual(result["browser_status"], "BLOCKED")
            self.assertFalse((Path(td) / "report.audit.png").exists())

    def test_29_text_extractor_emits_source_line_locator(self):
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "x.md"; p.write_text("# Heading\nBody", encoding="utf-8")
            self.assertIn("line[000002]", extract(p)[1].locator)

    def test_30_visual_review_required_until_reviewer_named(self):
        with tempfile.TemporaryDirectory() as td:
            src = Path(td) / "render"; out = Path(td) / "out"; src.mkdir(); Image.new("RGB", (1200, 675), "#cccccc").save(src / "p.png")
            r = audit_rendered_pages(src, out)
            self.assertEqual(r["pages"][0]["status"], "REVIEW_REQUIRED")

    def test_31_blank_visual_page_auto_fails(self):
        with tempfile.TemporaryDirectory() as td:
            src = Path(td) / "render"; src.mkdir(); Image.new("RGB", (1200, 675), "white").save(src / "p.png")
            ImageDraw = __import__("PIL.ImageDraw", fromlist=["ImageDraw"]).Draw
            im = Image.open(src / "p.png"); d = ImageDraw(im); d.rectangle((0,0,1199,674), fill="white"); im.save(src / "p.png")
            r = audit_rendered_pages(src, Path(td) / "out")
            self.assertEqual(r["pages"][0]["status"], "AUTO_FAIL")

    def test_32_explicit_visual_approval_requires_reviewer_and_timestamp(self):
        with tempfile.TemporaryDirectory() as td:
            src = Path(td) / "render"; src.mkdir(); Image.new("RGB", (1200, 675), "#cccccc").save(src / "p.png")
            review = Path(td) / "review.json"; review.write_text(json.dumps({"pages":[{"page":1,"reviewer":"R","reviewed_at":"2026-10-26","approved":True}]}), encoding="utf-8")
            r = audit_rendered_pages(src, Path(td) / "out", review_file=review)
            self.assertEqual(r["pages"][0]["status"], "HUMAN_APPROVED")


if __name__ == "__main__": unittest.main()
