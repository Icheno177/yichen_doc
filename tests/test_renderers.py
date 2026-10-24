from pathlib import Path
import tempfile
import json
import shutil
import unittest
from openpyxl import load_workbook
from pptx import Presentation
from docx import Document
from PIL import Image

from yichen_doc.cli import build
from yichen_doc.qa import inspect_output

ROOT = Path(__file__).resolve().parents[1]


class RendererAcceptanceTests(unittest.TestCase):
    def test_four_format_build_preserves_editable_objects(self):
        with tempfile.TemporaryDirectory() as tmp:
            output_dir = Path(tmp) / "deliverables"
            source = json.loads((ROOT / "examples" / "sample_research.json").read_text(encoding="utf-8"))
            source["sections"][0]["blocks"].append({"type":"chart","chart_type":"scatter","label":"Input relationship","categories":["A","B","C"],"x_values":[1,2,3],"x_unit":"x","unit":"y","series":[{"name":"Series 1","values":[2,4,3]}]})
            Image.new("RGB", (64, 32), (169, 35, 50)).save(Path(tmp) / "sample.png")
            source["sections"][0]["blocks"].append({"type":"image","image_path":"sample.png","alt":"Generated test swatch"})
            content_file = Path(tmp) / "content.json"; content_file.write_text(json.dumps(source, ensure_ascii=False), encoding="utf-8")
            files = build(content_file, output_dir, ["docx", "pptx", "xlsx", "html"])
            self.assertEqual({p.suffix for p in files}, {".docx", ".pptx", ".xlsx", ".html"})
            for p in files: inspect_output(p)
            doc = Document(next(p for p in files if p.suffix == ".docx"))
            self.assertGreaterEqual(len(doc.tables), 1)
            ppt = Presentation(next(p for p in files if p.suffix == ".pptx"))
            self.assertGreaterEqual(len(ppt.slides), 7)
            self.assertTrue(any(sh.has_chart for sl in ppt.slides for sh in sl.shapes))
            self.assertTrue(any(sh.has_table for sl in ppt.slides for sh in sl.shapes))
            self.assertTrue(any(sh.shape_type == 13 for sl in ppt.slides for sh in sl.shapes))
            wb = load_workbook(next(p for p in files if p.suffix == ".xlsx"), data_only=False)
            self.assertEqual(wb.sheetnames, ["README", "Sources", "Inputs", "Calculations", "Scenarios", "Sensitivity", "Dashboard", "QA"])
            self.assertTrue(any(isinstance(cell.value, str) and cell.value.startswith("=") for row in wb["Calculations"].iter_rows() for cell in row))
            self.assertGreater(len(wb["Dashboard"]._charts), 0)
            scatter = next(chart for chart in wb["Dashboard"]._charts if chart.__class__.__name__ == "ScatterChart")
            self.assertEqual(len(scatter.series), 1); self.assertIsNotNone(scatter.series[0].xVal)
            metrics = {m["id"]: m["value"] for m in source["metrics"]}
            self.assertEqual(metrics["M-01"] - metrics["M-02"] - metrics["M-03"], 550)
            page = next(p for p in files if p.suffix == ".html").read_text(encoding="utf-8")
            self.assertIn("tableSearch", page); self.assertIn("S-01", page); self.assertIn("data:image/png;base64,", page); self.assertIn("themeToggle", page); self.assertIn("localeCompare", page)


if __name__ == "__main__":
    unittest.main()
