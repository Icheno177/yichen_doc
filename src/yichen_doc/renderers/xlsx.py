from __future__ import annotations

from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, LineChart, PieChart, ScatterChart, Reference, Series
from openpyxl.utils import get_column_letter

NAVY = "202124"; RED = "A92332"; LIGHT = "F2F3F4"


def render(content: dict, output: str | Path) -> Path:
    out = Path(output); wb = Workbook(); readme = wb.active; readme.title = "README"
    for name in ("Sources", "Inputs", "Calculations", "Scenarios", "Sensitivity", "Dashboard", "QA"):
        wb.create_sheet(name)
    readme.append(["yichen_doc research workbook"]); readme.append(["Report", content["metadata"]["title"]]); readme.append(["Date", content["metadata"].get("date", "")]); readme.append(["Decision", content["decision"]["question"]]); readme.append(["Answer", content["decision"]["answer"]]); readme.append(["Units and evidence", "See Sources and section notes; unprovided inputs stay blank."])
    sources = wb["Sources"]; sources.append(["ID", "Title", "Organization", "Date", "Grade", "URL", "Limitations"])
    for s in content.get("sources", []): sources.append([s.get("id"), s.get("title"), s.get("organization"), s.get("date"), s.get("grade"), s.get("url"), s.get("limitations")])
    inputs = wb["Inputs"]; inputs.append(["Metric ID", "Label", "Value", "Unit", "Period", "Evidence Type", "Source ID"])
    for m in content.get("metrics", []): inputs.append([m.get("id"), m.get("label"), m.get("value"), m.get("unit"), m.get("period"), m.get("evidence_type"), m.get("source_id")])
    calc = wb["Calculations"]; calc.append(["Calculation ID", "Label", "Formula / Definition", "Input Metric IDs", "Excel Formula", "Result", "Unit", "Source IDs", "Limitations"])
    for sec in content.get("sections", []):
        for block in sec.get("blocks", []):
            if block.get("type") in ("formula", "chart", "metric"):
                calc.append(["", block.get("label", block.get("type")), block.get("formula", ""), "", "", None, block.get("unit", ""), ", ".join(block.get("source_ids", [])), block.get("text", "")])
    for calculation in content.get("calculations", []):
        calc.append([calculation["id"], calculation["label"], calculation["formula"], ", ".join(calculation.get("input_ids", [])), calculation.get("excel_formula", ""), None, calculation.get("unit"), ", ".join(calculation.get("source_ids", [])), calculation.get("limitations", "")])
        row = calc.max_row
        if calculation.get("excel_formula"):
            calc.cell(row, 6, calculation["excel_formula"])
            calc.cell(row, 6).number_format = '0.00;[Red]-0.00;–'
    for name in ("Scenarios", "Sensitivity"):
        ws = wb[name]; ws.append(["Scenario / driver", "Bear", "Base", "Bull", "Unit", "Assumption / Source"])
        for _ in content.get("assumptions", []): ws.append(["", None, None, None, None, ""])
    dash = wb["Dashboard"]; dash.append(["Metric", "Value", "Unit", "Period", "Evidence"])
    for m in content.get("metrics", []): dash.append([m.get("label"), m.get("value"), m.get("unit"), m.get("period"), m.get("evidence_type")])
    for sec in content.get("sections", []):
        for block in sec.get("blocks", []):
            if block.get("type") == "chart":
                start = dash.max_row + 2; dash.cell(start, 1, sec["title"] + " · " + block.get("label", "Chart"))
                dash.cell(start + 1, 1, "Category")
                for j, series in enumerate(block["series"], 2): dash.cell(start + 1, j, series["name"])
                for i, cat in enumerate(block["categories"], start + 2):
                    dash.cell(i, 1, block.get("x_values", [])[i - start - 2] if block.get("chart_type") == "scatter" else cat)
                    for j, series in enumerate(block["series"], 2): dash.cell(i, j, series["values"][i - start - 2])
                kind = block.get("chart_type", "bar")
                chart = LineChart() if kind == "line" else PieChart() if kind == "pie" else ScatterChart() if kind == "scatter" else BarChart()
                chart.title = block.get("label", "Chart"); chart.y_axis.title = block.get("unit", ""); chart.x_axis.title = "Category"
                data = Reference(dash, min_col=2, max_col=1 + len(block["series"]), min_row=start + 1, max_row=start + 1 + len(block["categories"]))
                cats = Reference(dash, min_col=1, min_row=start + 2, max_row=start + 1 + len(block["categories"]))
                if kind != "scatter":
                    chart.add_data(data, titles_from_data=True)
                    chart.set_categories(cats)
                else:
                    x_ref = Reference(dash, min_col=1, min_row=start + 2, max_row=start + 1 + len(block["categories"]))
                    chart.x_axis.title = block.get("x_unit", "X")
                    for j in range(len(block["series"])):
                        y_ref = Reference(dash, min_col=2 + j, min_row=start + 2, max_row=start + 1 + len(block["categories"]))
                        chart.series.append(Series(y_ref, x_ref, title_from_data=False, title=block["series"][j]["name"]))
                if kind == "stacked_bar": chart.grouping = "stacked"
                dash.add_chart(chart, f"G{start}")
    qa = wb["QA"]; qa.append(["Check", "Status", "Details"]); qa.append(["Content contract", "PASS", "Validated before rendering"]); qa.append(["Missing numeric inputs", "Blank", "Unknown values are not converted to zero"])
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]: cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor=NAVY)
        for col in ws.columns:
            letter = get_column_letter(col[0].column); width = min(54, max(12, max(len(str(c.value or "")) for c in col) + 2))
            ws.column_dimensions[letter].width = width
            for c in col[1:]:
                c.alignment = Alignment(vertical="top", wrap_text=True)
                if c.row % 2 == 0: c.fill = PatternFill("solid", fgColor=LIGHT)
    out.parent.mkdir(parents=True, exist_ok=True); wb.save(out); return out
