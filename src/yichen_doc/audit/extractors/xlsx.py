from __future__ import annotations

from pathlib import Path
from openpyxl import load_workbook
from ._common import add


def extract_xlsx(path: Path):
    wb = load_workbook(path, data_only=False, read_only=False)
    units = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                kind = "formula" if isinstance(cell.value, str) and cell.value.startswith("=") else "cell"
                item = add(path, f"sheet[{ws.title}]/cell[{cell.coordinate}]", kind, str(cell.value), ws.title)
                if item: units.append(item)
        for ti, table in enumerate(ws.tables.values(), 1):
            item = add(path, f"sheet[{ws.title}]/table[{ti}]/ref", "table", table.ref, ws.title)
            if item: units.append(item)
        for ci, chart in enumerate(ws._charts, 1):
            title = ""
            try:
                title = chart.title.tx.rich.p[0].r[0].t
            except Exception:
                pass
            item = add(path, f"sheet[{ws.title}]/chart[{ci}]/title", "chart", title or chart.__class__.__name__, ws.title)
            if item: units.append(item)
    return units
