from __future__ import annotations

from pathlib import Path
from pptx import Presentation
from ._common import add


def extract_pptx(path: Path):
    prs = Presentation(path)
    units = []
    for si, slide in enumerate(prs.slides, 1):
        for hi, shape in enumerate(slide.shapes, 1):
            loc = f"slide[{si:04d}]/shape[{hi:04d}]"
            if shape.has_text_frame and shape.text.strip():
                item = add(path, loc, "text", shape.text, f"slide {si}")
                if item: units.append(item)
            if shape.has_table:
                for ri, row in enumerate(shape.table.rows, 1):
                    for ci, cell in enumerate(row.cells, 1):
                        item = add(path, f"{loc}/table/row[{ri:04d}]/cell[{ci:03d}]", "table_cell", cell.text, f"slide {si}")
                        if item: units.append(item)
            if getattr(shape, "has_chart", False):
                chart = shape.chart
                categories = []
                try:
                    categories = [str(c.label) for c in chart.plots[0].categories]
                except Exception:
                    pass
                series_rows = []
                for series in chart.series:
                    series_rows.append({"name": series.name, "values": list(series.values)})
                import json
                item = add(path, loc + "/chart", "chart_data",
                           json.dumps({"chart_type": str(chart.chart_type), "categories": categories, "series": series_rows}, ensure_ascii=False), f"slide {si}")
                if item: units.append(item)
            if shape.shape_type == 13:
                item = add(path, loc + "/image", "image_resource", f"{shape.image.ext}:{shape.image.size[0]}x{shape.image.size[1]}", f"slide {si}")
                if item: units.append(item)
    return units
