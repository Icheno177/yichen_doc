"""Format-specific structural checks; rendering remains a separate check."""
from __future__ import annotations

from html.parser import HTMLParser
import json
from pathlib import Path
from urllib.parse import urlparse


def audit_docx(path: str | Path, golden_set: str | Path | None = None) -> dict:
    from docx import Document
    doc = Document(path)
    sections = []
    for s in doc.sections:
        sections.append({"page_width_in": round(s.page_width / 914400, 3), "page_height_in": round(s.page_height / 914400, 3),
                         "orientation": str(s.orientation), "different_first_page_header_footer": s.different_first_page_header_footer,
                         "margins_in": {k: round(getattr(s, f"{k}_margin") / 914400, 3) for k in ("top", "bottom", "left", "right")}})
    styles = {}
    style_properties = {}
    for style in doc.styles:
        if style.type != 1:  # paragraph styles
            continue
        font, pf = style.font, style.paragraph_format
        rpr = getattr(style._element, "rPr", None)
        rfonts = getattr(rpr, "rFonts", None) if rpr is not None else None
        east_asia = rfonts.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia") if rfonts is not None else None
        style_properties[style.name] = {"font_latin": font.name, "font_size_pt": float(font.size.pt) if font.size else None,
            "bold": font.bold, "italic": font.italic, "color": str(font.color.rgb) if font.color and font.color.rgb else None,
            "east_asia_font": east_asia, "space_before_pt": float(pf.space_before.pt) if pf.space_before else None,
            "space_after_pt": float(pf.space_after.pt) if pf.space_after else None,
            "line_spacing": float(pf.line_spacing) if isinstance(pf.line_spacing, (int, float)) else None,
            "first_line_indent_pt": float(pf.first_line_indent.pt) if pf.first_line_indent else None}
    for p in doc.paragraphs:
        if p.text.strip(): styles[p.style.name] = styles.get(p.style.name, 0) + 1
    header_footer_xml = "".join(s.header._element.xml + s.footer._element.xml for s in doc.sections)
    result = {"structure_status": "PASS", "render_status": "NOT_RUN", "sections": sections, "paragraph_style_usage": styles,
              "style_properties": style_properties, "table_count": len(doc.tables),
              "tables": [{"style": t.style.name if t.style else "", "rows": len(t.rows), "columns": len(t.columns)} for t in doc.tables],
              "paragraph_count": len(doc.paragraphs), "pictures_or_drawings": sum(1 for p in doc.paragraphs if "w:drawing" in p._p.xml),
              "page_number_field_present": "PAGE" in header_footer_xml,
              "headers_and_footers_present": any(s.header.paragraphs[0].text.strip() or s.footer.paragraphs[0].text.strip() for s in doc.sections)}
    if golden_set and Path(golden_set).is_file():
        golden = json.loads(Path(golden_set).read_text(encoding="utf-8"))
        expected = golden.get("document", {}).get("sections", [])
        expected_styles = golden.get("style_definitions", {})
        result["golden_set"] = str(Path(golden_set).resolve())
        result["golden_style_structure_comparison"] = {"expected_section_count": len(expected), "actual_section_count": len(sections), "margin_delta_in": [
            {k: round(sections[0]["margins_in"][k] - expected[0].get("margins_in", {}).get(k, sections[0]["margins_in"][k]), 3) for k in sections[0]["margins_in"]}
        ] if expected and sections else [], "missing_styles": sorted(set(expected_styles) - set(style_properties)),
            "style_property_deltas": {name: {key: {"expected": expected_styles[name].get(key), "actual": style_properties.get(name, {}).get(key)}
                                              for key in expected_styles[name] if expected_styles[name].get(key) != style_properties.get(name, {}).get(key)}
                                      for name in expected_styles if name in style_properties}}
    result["limitations"] = ["Style and section comparisons are structural checks; pagination, glyph fallback, clipping and visual fidelity require a true Word render.", "Direct run formatting and table border/cell-margin details are not fully compared against a Golden Set."]
    return result


def audit_xlsx(path: str | Path) -> dict:
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=False, read_only=False)
    sheets, formulas, formula_errors, external = [], [], [], []
    for ws in wb.worksheets:
        formula_count = 0; errors = []
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    formula_count += 1; formulas.append({"sheet": ws.title, "cell": c.coordinate, "formula": c.value, "recalculation": "NOT_RUN"})
                elif c.data_type == "e" or (isinstance(c.value, str) and c.value.startswith("#REF!")):
                    errors.append({"cell": c.coordinate, "value": c.value}); formula_errors.append({"sheet": ws.title, **errors[-1]})
        sheets.append({"name": ws.title, "state": ws.sheet_state, "max_row": ws.max_row, "max_column": ws.max_column,
                       "formula_count": formula_count, "data_validation_count": len(ws.data_validations.dataValidation),
                       "table_count": len(ws.tables), "chart_count": len(ws._charts), "error_cells": errors})
    for link in getattr(wb, "_external_links", []): external.append(str(link))
    return {"structure_status": "FAIL" if formula_errors else "PASS", "calculation_status": "NOT_RUN", "sheets": sheets,
            "formula_count": len(formulas), "formulas": formulas, "formula_error_cells": formula_errors,
            "external_links": external, "limitations": ["openpyxl reads formula definitions but does not calculate cached results.", "External sources and formula outputs require independent recalculation or a compatible spreadsheet engine.", "Hidden sheets are recorded and require reviewer interpretation; hidden content is not assumed malicious."]}


class _HTMLAudit(HTMLParser):
    def __init__(self):
        super().__init__(); self.ids = []; self.links = []; self.alt_missing = []; self.title_seen = False; self.errors = []
    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if a.get("id"): self.ids.append(a["id"])
        if tag == "a" and a.get("href"): self.links.append(("href", a["href"]))
        if tag in {"img", "script", "link", "iframe", "source"} and a.get("src"):
            self.links.append(("src", a["src"]))
        if tag == "img" and "alt" not in a: self.alt_missing.append(a.get("src", ""))
        if tag == "title": self.title_seen = True
    def error(self, message): self.errors.append(message)


def audit_html(path: str | Path, *, browser_check: bool = False, screenshot_path: str | Path | None = None) -> dict:
    p = Path(path); parser = _HTMLAudit()
    text = p.read_text(encoding="utf-8", errors="replace")
    parser.feed(text); parser.close()
    duplicates = sorted({x for x in parser.ids if parser.ids.count(x) > 1})
    resources = []
    for kind, raw in parser.links:
        u = urlparse(raw)
        if u.scheme in {"http", "https", "data", "mailto", "javascript"} or raw.startswith("#"):
            status = "EXTERNAL_OR_INLINE_UNVERIFIED" if u.scheme in {"http", "https"} else "NOT_LOCAL"
        else:
            target = (p.parent / raw.split("#", 1)[0].split("?", 1)[0]).resolve()
            status = "PRESENT" if target.is_file() else "MISSING"
        resources.append({"kind": kind, "reference": raw, "status": status})
    result = {"structure_status": "FAIL" if duplicates or any(x["status"] == "MISSING" for x in resources) else "PASS",
              "browser_status": "NOT_RUN", "title_tag_seen": parser.title_seen, "duplicate_ids": duplicates,
              "missing_alt_sources": parser.alt_missing, "resources": resources,
              "script_tag_count": text.lower().count("<script"), "console_errors": "NOT_OBSERVED"}
    if browser_check:
        if not screenshot_path:
            result["browser_status"] = "BLOCKED"
            result["browser_error"] = "Provide a separate screenshot_path; the source HTML is never modified."
            result["limitations"] = ["Browser screenshot output requires a separate destination path.", "A local-resource check or page load does not prove that linked sources support displayed claims.", "Browser rendering requires the optional Playwright package and an installed Chromium browser."]
            return result
        try:
            from playwright.sync_api import sync_playwright
            console_errors, page_errors = [], []
            with sync_playwright() as pw:
                browser = pw.chromium.launch(headless=True)
                page = browser.new_page()
                page.on("console", lambda msg: console_errors.append(msg.text) if msg.type == "error" else None)
                page.on("pageerror", lambda exc: page_errors.append(str(exc)))
                page.goto(p.resolve().as_uri(), wait_until="networkidle", timeout=30000)
                screenshot = Path(screenshot_path).resolve()
                screenshot.parent.mkdir(parents=True, exist_ok=True)
                page.screenshot(path=str(screenshot), full_page=True)
                result["browser_title"] = page.title()
                result["browser_screenshot"] = str(screenshot)
                browser.close()
            result["console_errors"] = console_errors; result["page_errors"] = page_errors
            result["browser_status"] = "PASS" if not console_errors and not page_errors else "AUTO_WARNING"
        except Exception as exc:
            result["browser_status"] = "BLOCKED"
            result["browser_error"] = f"{type(exc).__name__}: {exc}"
    result["limitations"] = ["A local-resource check or page load does not prove that linked sources support displayed claims.", "Browser rendering requires the optional Playwright package and an installed Chromium browser."]
    return result
